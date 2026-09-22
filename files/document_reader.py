"""
RENIX Document Reader
=====================

Unified document-reading layer for RENIX.

Responsibilities:
- Read plain-text documents
- Read Markdown
- Read JSON
- Read CSV
- Read XML
- Read HTML
- Read DOCX when python-docx is available
- Read PDF when PyMuPDF is available
- Detect document type automatically
- Extract searchable text
- Return structured document metadata
- Provide safe size limits
- Provide encoding detection/fallbacks

This module intentionally focuses on READING.
Editing/writing documents belongs to the appropriate file/coding
services.
"""

from __future__ import annotations

import csv
import io
import json
import logging
import mimetypes
import os
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator, Optional

logger = logging.getLogger(__name__)


# ============================================================
# EXCEPTIONS
# ============================================================


class DocumentReaderError(Exception):
    """Base exception for document-reader failures."""


class UnsupportedDocumentError(DocumentReaderError):
    """Raised when the document format is unsupported."""


class DocumentTooLargeError(DocumentReaderError):
    """Raised when a document exceeds the configured size limit."""


class DocumentEncodingError(DocumentReaderError):
    """Raised when a document cannot be decoded."""


class DocumentSecurityError(DocumentReaderError):
    """Raised when a document violates safety constraints."""


# ============================================================
# DATA MODELS
# ============================================================


@dataclass
class DocumentMetadata:
    path: str
    filename: str
    extension: str
    document_type: str
    mime_type: str
    size_bytes: int
    encoding: Optional[str] = None
    page_count: Optional[int] = None
    line_count: Optional[int] = None
    word_count: Optional[int] = None
    character_count: Optional[int] = None

    def to_dict(self) -> dict[str, Any]:
        return {
            "path": self.path,
            "filename": self.filename,
            "extension": self.extension,
            "document_type": self.document_type,
            "mime_type": self.mime_type,
            "size_bytes": self.size_bytes,
            "encoding": self.encoding,
            "page_count": self.page_count,
            "line_count": self.line_count,
            "word_count": self.word_count,
            "character_count": self.character_count,
        }


@dataclass
class DocumentPage:
    page_number: int
    text: str
    metadata: dict[str, Any] = field(
        default_factory=dict
    )

    def to_dict(self) -> dict[str, Any]:
        return {
            "page_number": self.page_number,
            "text": self.text,
            "metadata": self.metadata,
        }


@dataclass
class DocumentResult:
    success: bool
    text: str = ""
    metadata: Optional[DocumentMetadata] = None
    pages: list[DocumentPage] = field(
        default_factory=list
    )
    data: Any = None
    error: Optional[str] = None
    warnings: list[str] = field(
        default_factory=list
    )

    def to_dict(self) -> dict[str, Any]:
        return {
            "success": self.success,
            "text": self.text,
            "metadata": (
                self.metadata.to_dict()
                if self.metadata
                else None
            ),
            "pages": [
                page.to_dict()
                for page in self.pages
            ],
            "data": self.data,
            "error": self.error,
            "warnings": self.warnings,
        }


# ============================================================
# DOCUMENT READER
# ============================================================


class DocumentReader:
    """
    Unified document reader for RENIX.

    External libraries are optional. Core text formats work with
    the Python standard library.
    """

    DEFAULT_MAX_SIZE = 100 * 1024 * 1024

    TEXT_EXTENSIONS = {
        ".txt": "text",
        ".text": "text",
        ".md": "markdown",
        ".markdown": "markdown",
        ".rst": "text",
        ".log": "text",
        ".ini": "text",
        ".cfg": "text",
        ".conf": "text",
        ".yaml": "yaml",
        ".yml": "yaml",
    }

    STRUCTURED_EXTENSIONS = {
        ".json": "json",
        ".jsonl": "jsonl",
        ".csv": "csv",
        ".tsv": "tsv",
        ".xml": "xml",
    }

    WEB_EXTENSIONS = {
        ".html": "html",
        ".htm": "html",
    }

    PDF_EXTENSIONS = {
        ".pdf",
    }

    DOCX_EXTENSIONS = {
        ".docx",
    }

    def __init__(
        self,
        *,
        enabled: bool = True,
        max_file_size: int = DEFAULT_MAX_SIZE,
        default_encoding: str = "utf-8",
        preserve_newlines: bool = True,
    ) -> None:

        self.enabled = enabled
        self.max_file_size = max_file_size
        self.default_encoding = default_encoding
        self.preserve_newlines = preserve_newlines

        logger.info(
            "RENIX DocumentReader initialized."
        )

    # ========================================================
    # STATE
    # ========================================================

    def enable(self) -> None:
        self.enabled = True

    def disable(self) -> None:
        self.enabled = False

    def _check_enabled(self) -> None:
        if not self.enabled:
            raise DocumentReaderError(
                "DocumentReader is disabled."
            )

    # ========================================================
    # PATH VALIDATION
    # ========================================================

    @staticmethod
    def normalize_path(
        path: str | os.PathLike[str],
    ) -> Path:

        if path is None:
            raise DocumentReaderError(
                "Document path cannot be None."
            )

        return Path(path).expanduser()

    def validate_file(
        self,
        path: Path,
    ) -> None:

        if not path.exists():
            raise FileNotFoundError(
                f"Document does not exist: {path}"
            )

        if not path.is_file():
            raise DocumentReaderError(
                f"Path is not a file: {path}"
            )

        try:
            size = path.stat().st_size
        except OSError as exc:
            raise DocumentReaderError(
                f"Unable to inspect document: {path}"
            ) from exc

        if size > self.max_file_size:
            raise DocumentTooLargeError(
                f"Document is {size} bytes, exceeding "
                f"the limit of {self.max_file_size} bytes."
            )

    # ========================================================
    # TYPE DETECTION
    # ========================================================

    def detect_type(
        self,
        path: str | os.PathLike[str],
    ) -> str:

        document_path = self.normalize_path(path)
        extension = document_path.suffix.lower()

        if extension in self.TEXT_EXTENSIONS:
            return self.TEXT_EXTENSIONS[extension]

        if extension in self.STRUCTURED_EXTENSIONS:
            return self.STRUCTURED_EXTENSIONS[
                extension
            ]

        if extension in self.WEB_EXTENSIONS:
            return "html"

        if extension in self.PDF_EXTENSIONS:
            return "pdf"

        if extension in self.DOCX_EXTENSIONS:
            return "docx"

        mime_type, _ = mimetypes.guess_type(
            str(document_path)
        )

        if mime_type:

            if mime_type.startswith("text/"):
                return "text"

            if mime_type == "application/json":
                return "json"

            if mime_type == "application/pdf":
                return "pdf"

            if (
                mime_type
                == "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ):
                return "docx"

        return "unknown"

    # ========================================================
    # ENCODING
    # ========================================================

    def _decode_bytes(
        self,
        data: bytes,
    ) -> tuple[str, str]:

        encodings = [
            "utf-8",
            "utf-8-sig",
            "utf-16",
            "utf-16-le",
            "utf-16-be",
            "cp1252",
            "latin-1",
        ]

        # BOM-aware attempts first.
        if data.startswith(b"\xef\xbb\xbf"):
            encodings.insert(0, "utf-8-sig")

        elif data.startswith(
            b"\xff\xfe"
        ):
            encodings.insert(0, "utf-16")

        elif data.startswith(
            b"\xfe\xff"
        ):
            encodings.insert(0, "utf-16")

        seen = set()

        for encoding in encodings:

            if encoding in seen:
                continue

            seen.add(encoding)

            try:

                return (
                    data.decode(encoding),
                    encoding,
                )

            except UnicodeDecodeError:
                continue

        raise DocumentEncodingError(
            "Unable to decode document using supported encodings."
        )

    def read_text_file(
        self,
        path: Path,
    ) -> tuple[str, str]:

        try:

            data = path.read_bytes()

        except OSError as exc:

            raise DocumentReaderError(
                f"Unable to read document: {path}"
            ) from exc

        return self._decode_bytes(data)

    # ========================================================
    # TEXT NORMALIZATION
    # ========================================================

    def normalize_text(
        self,
        text: str,
    ) -> str:

        if self.preserve_newlines:

            text = text.replace(
                "\r\n",
                "\n",
            ).replace(
                "\r",
                "\n",
            )

        else:

            text = re.sub(
                r"\s+",
                " ",
                text,
            )

        return text

    @staticmethod
    def count_words(
        text: str,
    ) -> int:

        return len(
            re.findall(
                r"\S+",
                text,
            )
        )

    # ========================================================
    # METADATA
    # ========================================================

    def build_metadata(
        self,
        path: Path,
        document_type: str,
        *,
        encoding: Optional[str] = None,
        text: Optional[str] = None,
        page_count: Optional[int] = None,
    ) -> DocumentMetadata:

        size = path.stat().st_size

        mime_type, _ = mimetypes.guess_type(
            str(path)
        )

        if mime_type is None:
            mime_type = (
                "application/octet-stream"
            )

        line_count = None
        word_count = None
        character_count = None

        if text is not None:

            line_count = (
                text.count("\n") + 1
                if text
                else 0
            )

            word_count = self.count_words(text)
            character_count = len(text)

        return DocumentMetadata(
            path=str(path.resolve()),
            filename=path.name,
            extension=path.suffix.lower(),
            document_type=document_type,
            mime_type=mime_type,
            size_bytes=size,
            encoding=encoding,
            page_count=page_count,
            line_count=line_count,
            word_count=word_count,
            character_count=character_count,
        )

    # ========================================================
    # TEXT / MARKDOWN
    # ========================================================

    def read_text(
        self,
        path: str | os.PathLike[str],
    ) -> DocumentResult:

        self._check_enabled()

        document_path = self.normalize_path(path)
        self.validate_file(document_path)

        try:

            text, encoding = (
                self.read_text_file(document_path)
            )

            text = self.normalize_text(text)

            metadata = self.build_metadata(
                document_path,
                self.detect_type(document_path),
                encoding=encoding,
                text=text,
            )

            return DocumentResult(
                success=True,
                text=text,
                metadata=metadata,
            )

        except Exception as exc:

            logger.exception(
                "Failed to read text document."
            )

            return DocumentResult(
                success=False,
                error=str(exc),
            )

    # ========================================================
    # JSON
    # ========================================================

    def read_json(
        self,
        path: str | os.PathLike[str],
    ) -> DocumentResult:

        self._check_enabled()

        document_path = self.normalize_path(path)
        self.validate_file(document_path)

        try:

            text, encoding = (
                self.read_text_file(document_path)
            )

            data = json.loads(text)

            normalized = self.normalize_text(text)

            metadata = self.build_metadata(
                document_path,
                "json",
                encoding=encoding,
                text=normalized,
            )

            return DocumentResult(
                success=True,
                text=normalized,
                metadata=metadata,
                data=data,
            )

        except json.JSONDecodeError as exc:

            return DocumentResult(
                success=False,
                error=(
                    f"Invalid JSON at line "
                    f"{exc.lineno}, column "
                    f"{exc.colno}: {exc.msg}"
                ),
            )

        except Exception as exc:

            return DocumentResult(
                success=False,
                error=str(exc),
            )

    # ========================================================
    # JSONL
    # ========================================================

    def read_jsonl(
        self,
        path: str | os.PathLike[str],
    ) -> DocumentResult:

        self._check_enabled()

        document_path = self.normalize_path(path)
        self.validate_file(document_path)

        try:

            text, encoding = (
                self.read_text_file(document_path)
            )

            records = []

            for line_number, line in enumerate(
                text.splitlines(),
                start=1,
            ):

                stripped = line.strip()

                if not stripped:
                    continue

                try:

                    records.append(
                        json.loads(stripped)
                    )

                except json.JSONDecodeError as exc:

                    return DocumentResult(
                        success=False,
                        error=(
                            f"Invalid JSONL on line "
                            f"{line_number}: {exc}"
                        ),
                    )

            metadata = self.build_metadata(
                document_path,
                "jsonl",
                encoding=encoding,
                text=text,
            )

            return DocumentResult(
                success=True,
                text=text,
                metadata=metadata,
                data=records,
            )

        except Exception as exc:

            return DocumentResult(
                success=False,
                error=str(exc),
            )

    # ========================================================
    # CSV / TSV
    # ========================================================

    def read_csv(
        self,
        path: str | os.PathLike[str],
        *,
        delimiter: Optional[str] = None,
        has_header: bool = True,
    ) -> DocumentResult:

        self._check_enabled()

        document_path = self.normalize_path(path)
        self.validate_file(document_path)

        try:

            text, encoding = (
                self.read_text_file(document_path)
            )

            if delimiter is None:

                if document_path.suffix.lower() == ".tsv":
                    delimiter = "\t"
                else:
                    try:

                        dialect = csv.Sniffer().sniff(
                            text[:8192]
                        )

                        delimiter = dialect.delimiter

                    except csv.Error:
                        delimiter = ","

            reader = csv.reader(
                io.StringIO(text),
                delimiter=delimiter,
            )

            rows = list(reader)

            headers = None
            data_rows = rows

            if has_header and rows:

                headers = rows[0]
                data_rows = rows[1:]

            records = data_rows

            if headers:

                records = [
                    {
                        headers[index]: (
                            row[index]
                            if index < len(row)
                            else ""
                        )
                        for index in range(
                            len(headers)
                        )
                    }
                    for row in data_rows
                ]

            metadata = self.build_metadata(
                document_path,
                "tsv"
                if delimiter == "\t"
                else "csv",
                encoding=encoding,
                text=text,
            )

            return DocumentResult(
                success=True,
                text=text,
                metadata=metadata,
                data={
                    "headers": headers,
                    "rows": records,
                    "delimiter": delimiter,
                },
            )

        except Exception as exc:

            return DocumentResult(
                success=False,
                error=str(exc),
            )

    # ========================================================
    # XML
    # ========================================================

    def read_xml(
        self,
        path: str | os.PathLike[str],
    ) -> DocumentResult:

        self._check_enabled()

        document_path = self.normalize_path(path)
        self.validate_file(document_path)

        try:

            text, encoding = (
                self.read_text_file(document_path)
            )

            root = ET.fromstring(text)

            extracted_text = "\n".join(
                part.strip()
                for part in root.itertext()
                if part.strip()
            )

            metadata = self.build_metadata(
                document_path,
                "xml",
                encoding=encoding,
                text=extracted_text,
            )

            return DocumentResult(
                success=True,
                text=extracted_text,
                metadata=metadata,
                data=root,
            )

        except ET.ParseError as exc:

            return DocumentResult(
                success=False,
                error=f"Invalid XML: {exc}",
            )

        except Exception as exc:

            return DocumentResult(
                success=False,
                error=str(exc),
            )

    # ========================================================
    # HTML
    # ========================================================

    def read_html(
        self,
        path: str | os.PathLike[str],
    ) -> DocumentResult:

        self._check_enabled()

        document_path = self.normalize_path(path)
        self.validate_file(document_path)

        try:

            text, encoding = (
                self.read_text_file(document_path)
            )

            # Try BeautifulSoup if installed.
            try:

                from bs4 import BeautifulSoup

                soup = BeautifulSoup(
                    text,
                    "html.parser",
                )

                for element in soup(
                    [
                        "script",
                        "style",
                        "noscript",
                        "template",
                    ]
                ):
                    element.decompose()

                extracted = soup.get_text(
                    separator="\n"
                )

            except ImportError:

                # Safe fallback without external dependency.
                extracted = re.sub(
                    r"<script\b[^>]*>.*?</script>",
                    "",
                    text,
                    flags=re.IGNORECASE
                    | re.DOTALL,
                )

                extracted = re.sub(
                    r"<style\b[^>]*>.*?</style>",
                    "",
                    extracted,
                    flags=re.IGNORECASE
                    | re.DOTALL,
                )

                extracted = re.sub(
                    r"<[^>]+>",
                    " ",
                    extracted,
                )

            extracted = re.sub(
                r"[ \t]+",
                " ",
                extracted,
            )

            extracted = re.sub(
                r"\n\s*\n+",
                "\n\n",
                extracted,
            ).strip()

            metadata = self.build_metadata(
                document_path,
                "html",
                encoding=encoding,
                text=extracted,
            )

            return DocumentResult(
                success=True,
                text=extracted,
                metadata=metadata,
            )

        except Exception as exc:

            return DocumentResult(
                success=False,
                error=str(exc),
            )

    # ========================================================
    # DOCX
    # ========================================================

    def read_docx(
        self,
        path: str | os.PathLike[str],
    ) -> DocumentResult:

        self._check_enabled()

        document_path = self.normalize_path(path)
        self.validate_file(document_path)

        try:

            from docx import Document

        except ImportError:

            return DocumentResult(
                success=False,
                error=(
                    "DOCX support requires python-docx. "
                    "Install it with: pip install python-docx"
                ),
            )

        try:

            document = Document(
                str(document_path)
            )

            paragraphs = []

            for paragraph in document.paragraphs:

                if paragraph.text.strip():
                    paragraphs.append(
                        paragraph.text
                    )

            # Extract table contents too.
            tables = []

            for table in document.tables:

                table_rows = []

                for row in table.rows:

                    table_rows.append(
                        [
                            cell.text
                            for cell in row.cells
                        ]
                    )

                tables.append(table_rows)

                for row in table_rows:

                    paragraphs.append(
                        " | ".join(row)
                    )

            text = "\n".join(paragraphs)

            metadata = self.build_metadata(
                document_path,
                "docx",
                text=text,
            )

            return DocumentResult(
                success=True,
                text=text,
                metadata=metadata,
                data={
                    "paragraphs": [
                        paragraph.text
                        for paragraph
                        in document.paragraphs
                    ],
                    "tables": tables,
                },
            )

        except Exception as exc:

            return DocumentResult(
                success=False,
                error=str(exc),
            )

    # ========================================================
    # PDF
    # ========================================================

    def read_pdf(
        self,
        path: str | os.PathLike[str],
    ) -> DocumentResult:

        self._check_enabled()

        document_path = self.normalize_path(path)
        self.validate_file(document_path)

        try:

            import fitz

        except ImportError:

            return DocumentResult(
                success=False,
                error=(
                    "PDF support requires PyMuPDF. "
                    "Install it with: pip install PyMuPDF"
                ),
            )

        pages: list[DocumentPage] = []

        try:

            pdf = fitz.open(
                str(document_path)
            )

            try:

                for index in range(pdf.page_count):

                    page = pdf.load_page(index)

                    text = page.get_text(
                        "text"
                    )

                    pages.append(
                        DocumentPage(
                            page_number=index + 1,
                            text=text,
                            metadata={
                                "width": page.rect.width,
                                "height": page.rect.height,
                            },
                        )
                    )

            finally:

                pdf.close()

            full_text = "\n\n".join(
                page.text
                for page in pages
            )

            metadata = self.build_metadata(
                document_path,
                "pdf",
                text=full_text,
                page_count=len(pages),
            )

            return DocumentResult(
                success=True,
                text=full_text,
                metadata=metadata,
                pages=pages,
            )

        except Exception as exc:

            return DocumentResult(
                success=False,
                error=str(exc),
            )

    # ========================================================
    # GENERIC READ
    # ========================================================

    def read(
        self,
        path: str | os.PathLike[str],
    ) -> DocumentResult:

        self._check_enabled()

        document_path = self.normalize_path(path)
        document_type = self.detect_type(
            document_path
        )

        readers = {
            "text": self.read_text,
            "markdown": self.read_text,
            "yaml": self.read_text,
            "json": self.read_json,
            "jsonl": self.read_jsonl,
            "csv": self.read_csv,
            "tsv": self.read_csv,
            "xml": self.read_xml,
            "html": self.read_html,
            "docx": self.read_docx,
            "pdf": self.read_pdf,
        }

        reader = readers.get(document_type)

        if reader is None:

            return DocumentResult(
                success=False,
                error=(
                    f"Unsupported document type: "
                    f"{document_type}"
                ),
            )

        return reader(document_path)

    # ========================================================
    # ITERATE PAGES
    # ========================================================

    def iter_pages(
        self,
        path: str | os.PathLike[str],
    ) -> Iterator[DocumentPage]:

        result = self.read_pdf(path)

        if not result.success:

            raise DocumentReaderError(
                result.error or "Unable to read PDF."
            )

        yield from result.pages

    # ========================================================
    # SEARCH DOCUMENT
    # ========================================================

    def search(
        self,
        path: str | os.PathLike[str],
        query: str,
        *,
        case_sensitive: bool = False,
    ) -> list[dict[str, Any]]:

        if not query:
            return []

        result = self.read(path)

        if not result.success:
            raise DocumentReaderError(
                result.error or "Document reading failed."
            )

        text = result.text

        if not case_sensitive:

            text_to_search = text.lower()
            query_to_search = query.lower()

        else:

            text_to_search = text
            query_to_search = query

        matches = []

        position = 0

        while True:

            position = text_to_search.find(
                query_to_search,
                position,
            )

            if position == -1:
                break

            line_number = (
                text[:position].count("\n") + 1
            )

            start = max(
                0,
                position - 100,
            )

            end = min(
                len(text),
                position + len(query) + 100,
            )

            matches.append(
                {
                    "line": line_number,
                    "position": position,
                    "context": text[start:end],
                }
            )

            position += max(
                1,
                len(query),
            )

        return matches

    # ========================================================
    # PREVIEW
    # ========================================================

    def preview(
        self,
        path: str | os.PathLike[str],
        *,
        max_characters: int = 5000,
    ) -> DocumentResult:

        if max_characters <= 0:
            raise ValueError(
                "max_characters must be positive."
            )

        result = self.read(path)

        if not result.success:
            return result

        if len(result.text) > max_characters:

            result.text = (
                result.text[:max_characters]
                + "\n\n[RENIX: Preview truncated]"
            )

        return result

    # ========================================================
    # STATUS
    # ========================================================

    def get_status(self) -> dict[str, Any]:

        return {
            "enabled": self.enabled,
            "max_file_size": self.max_file_size,
            "default_encoding": self.default_encoding,
            "supported_formats": [
                "txt",
                "md",
                "markdown",
                "rst",
                "log",
                "ini",
                "cfg",
                "conf",
                "yaml",
                "yml",
                "json",
                "jsonl",
                "csv",
                "tsv",
                "xml",
                "html",
                "htm",
                "pdf",
                "docx",
            ],
            "optional_dependencies": {
                "docx": "python-docx",
                "pdf": "PyMuPDF",
                "html_parser": "beautifulsoup4",
            },
        }

    # ========================================================
    # SHUTDOWN
    # ========================================================

    def shutdown(self) -> None:

        self.enabled = False

        logger.info(
            "RENIX DocumentReader shut down."
        )


# ============================================================
# DEFAULT INSTANCE
# ============================================================

_default_reader: Optional[
    DocumentReader
] = None


def get_document_reader() -> DocumentReader:

    global _default_reader

    if _default_reader is None:
        _default_reader = DocumentReader()

    return _default_reader


# ============================================================
# CONVENIENCE FUNCTIONS
# ============================================================


def read_document(
    path: str | os.PathLike[str],
) -> DocumentResult:

    return get_document_reader().read(path)


def read_text(
    path: str | os.PathLike[str],
) -> DocumentResult:

    return get_document_reader().read_text(path)


def read_pdf(
    path: str | os.PathLike[str],
) -> DocumentResult:

    return get_document_reader().read_pdf(path)


def read_docx(
    path: str | os.PathLike[str],
) -> DocumentResult:

    return get_document_reader().read_docx(path)


def search_document(
    path: str | os.PathLike[str],
    query: str,
) -> list[dict[str, Any]]:

    return get_document_reader().search(
        path,
        query,
    )


# ============================================================
# PUBLIC API
# ============================================================

__all__ = [
    "DocumentReaderError",
    "UnsupportedDocumentError",
    "DocumentTooLargeError",
    "DocumentEncodingError",
    "DocumentSecurityError",
    "DocumentMetadata",
    "DocumentPage",
    "DocumentResult",
    "DocumentReader",
    "get_document_reader",
    "read_document",
    "read_text",
    "read_pdf",
    "read_docx",
    "search_document",
]


